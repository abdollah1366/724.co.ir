import tempfile
from decimal import Decimal
from io import StringIO
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from apps.ehr import importers
from apps.ehr.models import Encounter, EncounterService, GenericDrug, RVUCoefficient, ServiceCode
from apps.inventory.tests.fixtures import drug
from apps.lab.tests.fixtures import staff

from .test_ehr import appointment, doctor

OFFICIAL_FILE = Path(__file__).resolve().parents[3] / "data" / "catalogs" / "service-codes-1404.xlsx"


def workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RVU 3-1404"
    ws.append([None])
    ws.append(["گروه تعرفه"])
    ws.append(["کدملی (Code)", "ویژگی کد", "شرح کد (Value)", "توضیحات", " کل", "حرفه‌ای", "فنی", "ارزش پایه بیهوشی"])
    ws.append([100005, None, "آسپیراسیون سوزنی", "توضیح", 3.8, 3.8, None, 0])
    ws.append([100017, "*", "ساب سیژن", None, 4, 3, 1, "0"])
    ws.append([None, None, None])
    visits = wb.create_sheet("ویزیت")
    visits.append(["کد", "ویژگی کد", "شرح کد", "توضیحات", " کل", "حرفه‌ای", "فنی", "ارزش پایه بیهوشی"])
    visits.append([970000, "#", "ویزیت پزشک عمومی در مراکز سرپایی", None, 1.8, 1.3, 0.5, 0])
    general = wb.create_sheet("کدینگ جامع")
    general.append(["در این صفحه ..."])
    general.append(["کدملی", "شرح کد "])
    general.append([903500, "بسته پرستاری"])
    nursing = wb.create_sheet("پرستاری 1404")
    nursing.append(["کد ملی", "عنوان بسته خدمات", "ارزش نسبی"])
    nursing.append([903500, "بسته پرستاری جراحی", 8.8])
    modifiers = wb.create_sheet("کد تعدیلی")
    modifiers.append(["کد تعدیلی", "عنوان", "شرح کد"])
    modifiers.append([20, "عمل مستقل", "..."])
    wb.save(path)


class ServiceCodeImportTests(TestCase):
    def test_import_sheets_values_and_reimport(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "codes.xlsx"
            workbook(path)
            stats = importers.import_service_codes(path)
            self.assertEqual(stats["created"], 4)
            visit = ServiceCode.objects.get(code="970000")
            self.assertEqual((visit.total_rvu, visit.professional_rvu, visit.technical_rvu), (Decimal("1.8"), Decimal("1.3"), Decimal("0.5")))
            self.assertTrue(ServiceCode.objects.get(code="100017").is_cosmetic)
            self.assertEqual(ServiceCode.objects.get(code="903500").total_rvu, Decimal("8.8"))
            self.assertFalse(ServiceCode.objects.filter(code="20").exists())
            self.assertEqual(importers.import_service_codes(path)["created"], 0)
            self.assertEqual(ServiceCode.objects.count(), 4)

    def test_official_1404_file(self):
        out = StringIO()
        call_command("import_catalogs", "--service-codes", str(OFFICIAL_FILE), stdout=out)
        self.assertGreater(ServiceCode.objects.count(), 5900)
        visit = ServiceCode.objects.get(code="970000")
        self.assertIn("ویزیت پزشک عمومی", visit.title)
        self.assertEqual(visit.total_rvu, Decimal("1.8"))
        self.assertEqual(ServiceCode.objects.get(code="100010").technical_rvu, Decimal("3.1"))

    def test_tariff_uses_yearly_coefficients(self):
        visit = ServiceCode.objects.create(code="970000", title="ویزیت", total_rvu=Decimal("1.8"),
                                           professional_rvu=Decimal("1.3"), technical_rvu=Decimal("0.5"))
        self.assertIsNone(visit.tariff("private"))
        RVUCoefficient.objects.create(year=1403, sector="private", component="professional", value=100)
        RVUCoefficient.objects.create(year=1404, sector="private", component="professional", value=200_000)
        RVUCoefficient.objects.create(year=1404, sector="private", component="technical", value=100_000)
        self.assertEqual(visit.tariff("private"), 1.3 * 200_000 + 0.5 * 100_000)
        self.assertEqual(visit.tariff("private", 1403), 180)
        RVUCoefficient.objects.create(year=1404, sector="public", component="total", value=10_000)
        only_total = ServiceCode.objects.create(code="903500", title="پرستاری", total_rvu=Decimal("8.8"))
        self.assertEqual(only_total.tariff("public"), 88_000)


class GenericImportTests(TestCase):
    def test_csv_import_sample_and_product_linking(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "generics.csv"
            path.write_text("کد ژنریک,Generic Name,نام فارسی,شکل دارویی,Strength,ATC,دسته\n"
                            "12345,Metformin,متفورمین,قرص,500 mg,A10BA02,نسخه‌ای\n"
                            "12346,Paracetamol,استامینوفن,قرص,500 mg,N02BE01,OTC\n", encoding="utf-8")
            call_command("import_catalogs", "--generics", str(path), "--sample-generics", stdout=StringIO())
        metformin = GenericDrug.objects.get(name="Metformin", strength="500 mg")
        self.assertEqual((metformin.code, metformin.needs_verification), ("12345", False))  # نمونه، رسمی را بازنویسی نمی‌کند
        self.assertEqual(GenericDrug.objects.get(name="Paracetamol").category, "otc")
        self.assertTrue(GenericDrug.objects.get(name="Alprazolam").needs_verification)
        product = drug("glucophage", generic_name="metformin", strength="500 mg")
        self.assertEqual(importers.link_products_to_generics(), 1)
        product.refresh_from_db()
        self.assertEqual(product.generic, metformin)

    def test_icd_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "icd.csv"
            path.write_text("code,title,title_fa\nJ00,Acute nasopharyngitis,سرماخوردگی\nbad,x,y\n", encoding="utf-8")
            call_command("import_catalogs", "--icd10", str(path), stdout=StringIO())
        from apps.ehr.models import ICD10Code

        self.assertEqual(ICD10Code.objects.get(code="J00").title_fa, "سرماخوردگی")


class EncounterServiceTests(TestCase):
    def test_doctor_adds_service_with_tariff_snapshot(self):
        user = staff("doc", "پزشک")
        appt = appointment(doctor(user))
        visit = ServiceCode.objects.create(code="970000", title="ویزیت پزشک عمومی", total_rvu=Decimal("1.8"),
                                           professional_rvu=Decimal("1.3"), technical_rvu=Decimal("0.5"))
        ServiceCode.objects.create(code="100017", attribute="*", title="ساب سیژن", total_rvu=4)
        RVUCoefficient.objects.create(year=1404, sector="private", component="professional", value=200_000)
        RVUCoefficient.objects.create(year=1404, sector="private", component="technical", value=100_000)
        self.client.force_login(user)
        self.client.post(reverse("ehr:encounter_start", args=[appt.patient.pk]), {"appointment": appt.pk})
        encounter = Encounter.objects.get()
        url = reverse("ehr:encounter_service", args=[encounter.pk])
        self.client.post(url, {"s-code": "۹۷۰۰۰۰", "s-quantity": 1, "s-sector": "private"})
        self.client.post(url, {"s-code": "999999", "s-quantity": 1, "s-sector": "private"})
        item = EncounterService.objects.get()
        self.assertEqual((item.service_code, item.unit_tariff), (visit, 310_000))
        RVUCoefficient.objects.filter(component="professional").update(value=1)
        item.refresh_from_db()
        self.assertEqual(item.unit_tariff, 310_000)  # تعرفه زمان ثبت حفظ می‌شود
        page = self.client.get(reverse("ehr:encounter", args=[encounter.pk]))
        self.assertContains(page, "970000")
        data = self.client.get(reverse("ehr:service_search"), {"q": "ساب"}).json()["results"]
        self.assertEqual(data[0]["cosmetic"], True)
